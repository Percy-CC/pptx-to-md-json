from pathlib import Path
import colorsys
import json
import re
from urllib.parse import quote

from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn

# --------------------------------------------------------------------------
# Configuration
# --------------------------------------------------------------------------
BASE_DIR = Path(__file__).resolve().parent
INPUT_DIR = BASE_DIR / "input"
OUTPUT_DIR = BASE_DIR / "output"

INCLUDE_COLORS = True                  # set False to drop all colour spans
IGNORE_COLORS = {"000000", "FFFFFF"}   # colours never worth reporting
REPEAT_MERGED_CELLS = True             # copy merged-cell text into every spanned cell
ROW_TOLERANCE_EMU = 91440              # 0.1 inch: shapes this close vertically = same row
LINE_BREAK_MD = "<br>"                 # how a soft line break (Shift+Enter) is written

SKIP_PH = {PP_PLACEHOLDER.FOOTER, PP_PLACEHOLDER.SLIDE_NUMBER, PP_PLACEHOLDER.DATE}
TITLE_PH = {PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE, PP_PLACEHOLDER.VERTICAL_TITLE}
OLE_TYPES = {MSO_SHAPE_TYPE.EMBEDDED_OLE_OBJECT, MSO_SHAPE_TYPE.LINKED_OLE_OBJECT}
TRUE_VALUES = {"1", "true", "on"}
DGM_NS = "http://schemas.openxmlformats.org/drawingml/2006/diagram"

TAG_R, TAG_FLD, TAG_BR = qn("a:r"), qn("a:fld"), qn("a:br")
TAG_RPR, TAG_T = qn("a:rPr"), qn("a:t")
TAG_PIC, TAG_GRP = qn("p:pic"), qn("p:grpSp")

# (bold, italic, underline, strike, baseline_sign, color, highlight, link)
EMPTY_STYLE = (False, False, False, False, 0, None, None, None)


# --------------------------------------------------------------------------
# Text helpers / escaping
# --------------------------------------------------------------------------
def clean(text):
    return " ".join((text or "").replace("\x0b", " ").split()).strip()


_MD_SPECIAL = re.compile(r"([\\`*_\[\]|~])")


def md_escape(s):
    s = _MD_SPECIAL.sub(r"\\\1", s)
    return s.replace("<", "&lt;").replace(">", "&gt;")


def escape_line_start(s):
    """Stop text that starts with '#', '-', '1.' etc. from becoming headings/lists."""
    m = re.match(r"^(\d+)([.)])", s)
    if m:
        return m.group(1) + "\\" + m.group(2) + s[m.end():]
    if s and s[0] in "#-+=":
        return "\\" + s
    return s


def md_url(url):
    return quote(url.strip(), safe=":/?#[]@!$&'*+,;=%")


# --------------------------------------------------------------------------
# Colours (explicit RGB + theme colours with lumMod/lumOff/tint/shade)
# --------------------------------------------------------------------------
def _localname(el):
    return etree.QName(el).localname if isinstance(el.tag, str) else None


def apply_color_mods(hex_val, color_el):
    try:
        r, g, b = (int(hex_val[i:i + 2], 16) / 255 for i in (0, 2, 4))
    except (ValueError, TypeError):
        return None
    for mod in color_el:
        name = _localname(mod)
        try:
            v = int(mod.get("val")) / 100000
        except (TypeError, ValueError):
            continue
        if name in ("lumMod", "lumOff"):
            h, l, s = colorsys.rgb_to_hls(r, g, b)
            l = l * v if name == "lumMod" else l + v
            r, g, b = colorsys.hls_to_rgb(h, min(max(l, 0), 1), s)
        elif name == "shade":
            r, g, b = r * v, g * v, b * v
        elif name == "tint":
            r, g, b = (c + (1 - c) * (1 - v) for c in (r, g, b))
    return "".join(f"{round(min(max(c, 0), 1) * 255):02X}" for c in (r, g, b))


def resolve_color(container, theme):
    """container = <a:solidFill> or <a:highlight>; returns 'RRGGBB' or None."""
    if container is None:
        return None
    for child in container:
        name = _localname(child)
        if name == "srgbClr":
            base = child.get("val")
        elif name == "schemeClr":
            base = theme.get(child.get("val"))
        elif name == "sysClr":
            base = child.get("lastClr")
        else:
            continue
        if base:
            return apply_color_mods(base.upper(), child)
    return None


def load_theme_colors(master):
    colors = {}
    try:
        theme_part = master.part.part_related_by(RT.THEME)
        root = etree.fromstring(theme_part.blob)
    except Exception:
        return colors
    scheme = root.find(".//" + qn("a:clrScheme"))
    if scheme is not None:
        for slot in scheme:
            slot_name = _localname(slot)
            if not slot_name:
                continue
            for c in slot:
                n = _localname(c)
                val = c.get("val") if n == "srgbClr" else c.get("lastClr") if n == "sysClr" else None
                if val:
                    colors[slot_name] = val.upper()
                    break
    clr_map = master._element.find(qn("p:clrMap"))
    mapping = dict(clr_map.attrib) if clr_map is not None else {
        "bg1": "lt1", "tx1": "dk1", "bg2": "lt2", "tx2": "dk2"}
    for alias, target in mapping.items():
        if target in colors:
            colors[alias] = colors[target]
    return colors


# --------------------------------------------------------------------------
# Style inheritance (run -> shape lstStyle -> layout -> master -> txStyles)
# --------------------------------------------------------------------------
class SlideEnv:
    def __init__(self, theme, master_el, prs_default, label):
        self.theme = theme
        self.master_el = master_el
        self.prs_default = prs_default
        self.label = label


class TextContext:
    def __init__(self, list_styles, theme, part):
        self.list_styles = [x for x in list_styles if x is not None]
        self.theme = theme
        self.part = part
        self._cache = {}

    def defaults(self, level):
        if level not in self._cache:
            tag = qn(f"a:lvl{min(level, 8) + 1}pPr")
            found = []
            for lst in self.list_styles:
                lvl = lst.find(tag)
                if lvl is not None:
                    d = lvl.find(qn("a:defRPr"))
                    if d is not None:
                        found.append(d)
            self._cache[level] = found
        return self._cache[level]


def list_styles_for(shape, env):
    styles = []
    lst_path = f"{qn('p:txBody')}/{qn('a:lstStyle')}"
    cur = shape
    for _ in range(3):  # slide shape -> layout placeholder -> master placeholder
        if cur is None:
            break
        styles.append(cur._element.find(lst_path))
        if not getattr(cur, "is_placeholder", False):
            break
        try:
            cur = cur._base_placeholder
        except Exception:
            cur = None

    if getattr(shape, "is_placeholder", False):
        tx = env.master_el.find(qn("p:txStyles")) if env.master_el is not None else None
        if tx is not None:
            try:
                ph_type = shape.placeholder_format.type
            except Exception:
                ph_type = None
            styles.append(tx.find(qn("p:titleStyle" if ph_type in TITLE_PH else "p:bodyStyle")))
    else:
        styles.append(env.prs_default)
    return styles


def make_ctx(shape, env):
    try:
        styles = list_styles_for(shape, env)
    except Exception:
        styles = []
    return TextContext(styles, env.theme, shape.part)


def hyperlink_of(rPr, part):
    if rPr is None or part is None:
        return None
    h = rPr.find(qn("a:hlinkClick"))
    if h is None:
        return None
    rid = h.get(qn("r:id"))
    if not rid:
        return None
    try:
        rel = part.rels[rid]
        return rel.target_ref if rel.is_external else None
    except Exception:
        return None


def run_style(rPr, defaults, ctx):
    chain = [el for el in [rPr] + defaults if el is not None]

    def attr(name):
        for el in chain:
            v = el.get(name)
            if v is not None:
                return v
        return None

    bold = attr("b") in TRUE_VALUES
    italic = attr("i") in TRUE_VALUES
    underline = attr("u") not in (None, "none")
    strike = attr("strike") in ("sngStrike", "dblStrike")

    baseline = 0
    raw = attr("baseline")
    if raw:
        try:
            b = float(raw.rstrip("%"))
            baseline = 1 if b > 0 else -1 if b < 0 else 0
        except ValueError:
            pass

    color = None
    if INCLUDE_COLORS and rPr is not None:
        explicit = resolve_color(rPr.find(qn("a:solidFill")), ctx.theme)
        if explicit:
            inherited = None
            for d in defaults:
                sf = d.find(qn("a:solidFill"))
                if sf is not None:
                    inherited = resolve_color(sf, ctx.theme)
                    break
            inherited = inherited or ctx.theme.get("tx1")
            if explicit != inherited and explicit not in IGNORE_COLORS:
                color = explicit

    highlight = None
    if rPr is not None:
        highlight = resolve_color(rPr.find(qn("a:highlight")), ctx.theme)

    link = hyperlink_of(rPr, ctx.part)
    return (bold, italic, underline, strike, baseline, color, highlight, link)


# --------------------------------------------------------------------------
# Paragraph -> Markdown
# --------------------------------------------------------------------------
def render_group(style, text):
    core = text.strip()
    if not core:
        return text
    lead = text[:len(text) - len(text.lstrip())]
    trail = text[len(text.rstrip()):]
    core = md_escape(core)
    bold, italic, underline, strike, baseline, color, highlight, link = style

    if bold and italic:
        core = f"***{core}***"
    elif bold:
        core = f"**{core}**"
    elif italic:
        core = f"*{core}*"
    if strike:
        core = f"~~{core}~~"
    if underline:
        core = f"<u>{core}</u>"
    if highlight:
        core = f'<mark style="background:#{highlight}">{core}</mark>'
    if color:
        core = f'<span style="color:#{color}">{core}</span>'
    if baseline > 0:
        core = f"<sup>{core}</sup>"
    elif baseline < 0:
        core = f"<sub>{core}</sub>"
    if link:
        core = f"[{core}]({md_url(link)})"
    return lead + core + trail


def paragraph_md(p, ctx):
    defaults = ctx.defaults(p.level)

    # 1. collect pieces: runs, fields, line breaks
    pieces = []
    for child in p._p.iterchildren():
        if child.tag in (TAG_R, TAG_FLD):
            text = (child.findtext(TAG_T) or "").replace("\x0b", " ")
            if not text:
                continue
            if not text.strip():
                pieces.append((None, text))           # whitespace-only run
            else:
                pieces.append((run_style(child.find(TAG_RPR), defaults, ctx), text))
        elif child.tag == TAG_BR:
            pieces.append(("BR", None))

    # 2. merge adjacent runs with identical formatting
    groups = []
    for style, text in pieces:
        if style == "BR":
            groups.append(["BR", ""])
        elif style is None:
            if groups and groups[-1][0] != "BR":
                groups[-1][1] += text
            else:
                groups.append([EMPTY_STYLE, text])
        elif groups and groups[-1][0] == style:
            groups[-1][1] += text
        else:
            groups.append([style, text])

    # 3. render
    out = [LINE_BREAK_MD if style == "BR" else render_group(style, text)
           for style, text in groups]
    md = re.sub(r"[ \t]+", " ", "".join(out)).strip()
    while md.startswith(LINE_BREAK_MD):
        md = md[len(LINE_BREAK_MD):].strip()
    while md.endswith(LINE_BREAK_MD):
        md = md[:-len(LINE_BREAK_MD)].strip()
    return escape_line_start(md)


def paragraphs_of(text_frame, ctx):
    items = []
    for p in text_frame.paragraphs:
        plain = clean(p.text)
        if not plain:
            continue
        try:
            md = paragraph_md(p, ctx)
        except Exception:
            md = ""
        items.append({"level": p.level, "text": plain,
                      "md": md or escape_line_start(md_escape(plain))})
    return items


# --------------------------------------------------------------------------
# Tables, charts, SmartArt
# --------------------------------------------------------------------------
def table_block(shape, env):
    table = shape.table
    ctx = TextContext([], env.theme, shape.part)
    md_grid, plain_grid = [], []
    for r in table.rows:
        md_row, plain_row = [], []
        for c in r.cells:
            paras = [paragraph_md(p, ctx) for p in c.text_frame.paragraphs]
            md_row.append("<br>".join(x for x in paras if x))
            plain_row.append(clean(" ".join(p.text for p in c.text_frame.paragraphs)))
        md_grid.append(md_row)
        plain_grid.append(plain_row)
    if not md_grid:
        return None

    if REPEAT_MERGED_CELLS:
        try:
            for ri, r in enumerate(table.rows):
                for ci, c in enumerate(r.cells):
                    if c.is_merge_origin:
                        for dr in range(c.span_height):
                            for dc in range(c.span_width):
                                rr, cc = ri + dr, ci + dc
                                if rr < len(md_grid) and cc < len(md_grid[rr]):
                                    md_grid[rr][cc] = md_grid[ri][ci]
                                    plain_grid[rr][cc] = plain_grid[ri][ci]
        except Exception:
            pass

    width = max(len(r) for r in md_grid)
    md_grid = [r + [""] * (width - len(r)) for r in md_grid]

    try:
        has_header = bool(table.first_row)
    except Exception:
        has_header = True
    if has_header:
        header, body = md_grid[0], md_grid[1:]
    else:
        header, body = [" "] * width, md_grid

    lines = ["| " + " | ".join(header) + " |",
             "| " + " | ".join(["---"] * width) + " |"]
    lines += ["| " + " | ".join(r) + " |" for r in body]
    return {"type": "table", "has_header": has_header,
            "rows": plain_grid, "markdown": "\n".join(lines)}


def _fmt_num(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v)


def chart_block(chart):
    title = ""
    try:
        if chart.has_title and chart.chart_title.has_text_frame:
            title = clean(chart.chart_title.text_frame.text)
    except Exception:
        pass

    tables = []
    try:
        for plot in chart.plots:
            series = list(plot.series)
            if not series:
                continue
            try:
                cats = [str(c) for c in plot.categories]
            except Exception:
                cats = []
            names = [s.name or f"Series {i + 1}" for i, s in enumerate(series)]
            values = []
            for s in series:
                try:
                    values.append(list(s.values))
                except Exception:
                    values.append([])
            n = max([len(cats)] + [len(v) for v in values])
            lines = ["| " + " | ".join(md_escape(x) for x in ["Category"] + names) + " |",
                     "| " + " | ".join(["---"] * (len(names) + 1)) + " |"]
            for i in range(n):
                cat = cats[i] if i < len(cats) else str(i + 1)
                row = [md_escape(cat)] + [_fmt_num(v[i]) if i < len(v) else "" for v in values]
                lines.append("| " + " | ".join(row) + " |")
            tables.append("\n".join(lines))
    except Exception:
        pass

    return {"type": "chart", "text": title or "(chart)", "markdown": "\n\n".join(tables)}


def is_smartart(shape):
    el = shape._element
    if el.tag != qn("p:graphicFrame"):
        return False
    uris = el.xpath("./a:graphic/a:graphicData/@uri")
    return bool(uris) and uris[0] == DGM_NS


def smartart_block(shape):
    rel_el = next(shape._element.iter(f"{{{DGM_NS}}}relIds"), None)
    if rel_el is None:
        return None
    rid = rel_el.get(qn("r:dm"))
    root = etree.fromstring(shape.part.related_part(rid).blob)

    parent = {}
    for cxn in root.iter(f"{{{DGM_NS}}}cxn"):
        if cxn.get("type", "parOf") == "parOf":
            parent[cxn.get("destId")] = cxn.get("srcId")

    def depth(model_id):
        d, seen = 0, set()
        while model_id in parent and model_id not in seen:
            seen.add(model_id)
            model_id = parent[model_id]
            d += 1
        return d

    paras = []
    for pt in root.iter(f"{{{DGM_NS}}}pt"):
        if pt.get("type", "node") not in ("node", "asst"):
            continue
        lines = ["".join(t.text or "" for t in ap.iter(TAG_T)) for ap in pt.iter(qn("a:p"))]
        text = clean(" ".join(lines))
        if text:
            paras.append({"level": max(depth(pt.get("modelId")) - 1, 0), "text": text,
                          "md": escape_line_start(md_escape(text))})
    return {"type": "smartart", "paragraphs": paras} if paras else None


# --------------------------------------------------------------------------
# Shapes
# --------------------------------------------------------------------------
def safe_shape_type(shape):
    try:
        return shape.shape_type
    except Exception:
        return None


def is_skipped(shape):
    if getattr(shape, "is_placeholder", False):
        try:
            return shape.placeholder_format.type in SKIP_PH
        except Exception:
            return False
    return False


def reading_order(shapes):
    """Group shapes into rows (tops within tolerance), then order each row left to right."""
    items = sorted(shapes, key=lambda s: (s.top or 0, s.left or 0))
    rows = []
    for s in items:
        top = s.top or 0
        if rows and top - rows[-1][0] <= ROW_TOLERANCE_EMU:
            rows[-1][1].append(s)
        else:
            rows.append([top, [s]])
    return [s for _, row in rows for s in sorted(row, key=lambda s: s.left or 0)]


def safe_extract(shape, env):
    try:
        return extract_shape(shape, env)
    except Exception as e:
        sid = getattr(shape, "shape_id", "?")
        print(f"Warning: {env.label}: skipped shape {sid}: {type(e).__name__}: {e}")
        return []


def extract_shape(shape, env):
    if is_skipped(shape):
        return []

    el = shape._element
    stype = safe_shape_type(shape)
    blocks = []

    if el.tag == TAG_GRP or stype == MSO_SHAPE_TYPE.GROUP:
        for child in reading_order(list(shape.shapes)):
            blocks.extend(safe_extract(child, env))
        return blocks

    if getattr(shape, "has_text_frame", False):
        paras = paragraphs_of(shape.text_frame, make_ctx(shape, env))
        if paras:
            blocks.append({"type": "text", "paragraphs": paras})

    if getattr(shape, "has_table", False):
        tb = table_block(shape, env)
        if tb:
            blocks.append(tb)

    if getattr(shape, "has_chart", False):
        blocks.append(chart_block(shape.chart))

    if is_smartart(shape):
        sb = smartart_block(shape)
        if sb:
            blocks.append(sb)

    if stype in OLE_TYPES:
        try:
            prog = shape.ole_format.prog_id or "unknown"
        except Exception:
            prog = "unknown"
        blocks.append({"type": "object", "text": f"Embedded object ({prog})"})

    if el.tag == TAG_PIC:  # also catches pictures inside placeholders
        alt = el.xpath("./p:nvPicPr/p:cNvPr/@descr")
        if alt and clean(alt[0]):
            blocks.append({"type": "image", "text": clean(alt[0])})

    return blocks


# --------------------------------------------------------------------------
# Slides
# --------------------------------------------------------------------------
def get_title(slide):
    t = slide.shapes.title
    return clean(t.text) if t is not None and t.has_text_frame else ""


def get_notes(slide):
    if not slide.has_notes_slide:
        return ""
    tf = slide.notes_slide.notes_text_frame
    if tf is None:
        return ""
    return "\n".join(clean(p.text) for p in tf.paragraphs if clean(p.text))


def is_hidden(slide):
    return slide._element.get("show") in ("0", "false")


def write_markdown(all_slides, path):
    with open(path, "w", encoding="utf-8") as f:
        f.write("# PPTX text corpus\n\n")
        current_file = None
        for s in all_slides:
            if s["source_file"] != current_file:
                current_file = s["source_file"]
                f.write(f"## File: {md_escape(current_file)}\n\n")

            title = md_escape(s["title"]) if s["title"] else "(No title)"
            hidden = " [hidden]" if s["hidden"] else ""
            f.write(f"### Slide {s['slide_number']}: {title}{hidden}\n\n")

            for b in s["blocks"]:
                if b["type"] in ("text", "smartart"):
                    min_lvl = min(p["level"] for p in b["paragraphs"])
                    for p in b["paragraphs"]:
                        f.write("  " * (p["level"] - min_lvl) + f"- {p['md']}\n")
                    f.write("\n")
                elif b["type"] == "table":
                    f.write(b["markdown"] + "\n\n")
                elif b["type"] == "chart":
                    f.write(f"*Chart: {md_escape(b['text'])}*\n\n")
                    if b.get("markdown"):
                        f.write(b["markdown"] + "\n\n")
                elif b["type"] == "image":
                    f.write(f"*Image alt text: {md_escape(b['text'])}*\n\n")
                elif b["type"] == "object":
                    f.write(f"*{md_escape(b['text'])}*\n\n")

            if s["speaker_notes"]:
                f.write("**Speaker notes**\n\n")
                for line in s["speaker_notes"].splitlines():
                    f.write(f"> {escape_line_start(md_escape(line))}\n")
                f.write("\n")


def main():
    if not INPUT_DIR.exists():
        raise SystemExit(f"Input folder not found: {INPUT_DIR}")

    files = [p for p in sorted(INPUT_DIR.iterdir())
             if p.is_file() and p.suffix.lower() == ".pptx" and not p.name.startswith("~$")]
    if not files:
        raise SystemExit(f"No .pptx files in: {INPUT_DIR}")

    OUTPUT_DIR.mkdir(exist_ok=True)

    all_slides = []
    for pptx_path in files:
        try:
            presentation = Presentation(pptx_path)
        except Exception as e:
            print(f"Skipped {pptx_path.name}: {e}")
            continue

        prs_default = presentation._element.find(qn("p:defaultTextStyle"))
        theme_cache = {}

        for number, slide in enumerate(presentation.slides, start=1):
            label = f"{pptx_path.name} slide {number}"
            try:
                master = slide.slide_layout.slide_master
                key = str(master.part.partname)
                if key not in theme_cache:
                    theme_cache[key] = load_theme_colors(master)
                env = SlideEnv(theme_cache[key], master._element, prs_default, label)
            except Exception:
                env = SlideEnv({}, None, prs_default, label)

            title_shape = slide.shapes.title
            title_id = title_shape.shape_id if title_shape is not None else None

            blocks = []
            for shape in reading_order(list(slide.shapes)):
                if title_id is not None and shape.shape_id == title_id:
                    continue
                blocks.extend(safe_extract(shape, env))

            try:
                notes = get_notes(slide)
            except Exception:
                notes = ""

            all_slides.append({
                "source_file": pptx_path.name,
                "slide_number": number,
                "hidden": is_hidden(slide),
                "title": get_title(slide),
                "blocks": blocks,
                "speaker_notes": notes,
            })

    with open(OUTPUT_DIR / "all_slides.json", "w", encoding="utf-8") as f:
        json.dump(all_slides, f, ensure_ascii=False, indent=2)

    write_markdown(all_slides, OUTPUT_DIR / "all_slides.md")

    print(f"Converted {len(all_slides)} slides.")
    print(f"Created: {OUTPUT_DIR / 'all_slides.md'}")
    print(f"Created: {OUTPUT_DIR / 'all_slides.json'}")


if __name__ == "__main__":
    main()